"""Citirea documentelor (PDF, Word, text), alegerea unei părți din ele și
căutarea după cuvinte cheie în documente lungi."""

import hashlib
import io
import math
import re
import unicodedata
from collections import Counter

import pymupdf
import streamlit as st

from config import (
    MAX_CARACTERE,
    NR_FRAGMENTE_INTREBARE,
    PLURAL,
    TINTA_SECTIUNE,
)
from utile import fmt


# ------------------------- citirea documentelor -------------------------
def imparte_text(text, tinta=TINTA_SECTIUNE):
    """Împarte un text fără pagini în secțiuni de ~tinta caractere,
    tăind la limite de paragraf sau de propoziție."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    blocuri = [b.strip() for b in re.split(r"\n\s*\n", text) if b.strip()]
    if len(blocuri) < 2:
        blocuri = [b.strip() for b in text.split("\n") if b.strip()]

    mici = []
    for b in blocuri:
        while len(b) > tinta:
            taie = b.rfind(". ", 0, tinta)
            if taie < tinta // 2:
                taie = b.rfind(" ", 0, tinta)
            if taie < tinta // 2:
                taie = tinta - 1
            mici.append(b[: taie + 1].strip())
            b = b[taie + 1 :].strip()
        if b:
            mici.append(b)

    sectiuni, curent = [], ""
    for b in mici:
        if curent and len(curent) + len(b) + 2 > tinta:
            sectiuni.append(curent)
            curent = b
        else:
            curent = f"{curent}\n\n{b}" if curent else b
    if curent:
        sectiuni.append(curent)
    return [{"pagina": i, "text": t} for i, t in enumerate(sectiuni, start=1)]


def citeste_txt(octeti):
    for enc in ("utf-8-sig", "cp1250", "latin-1"):
        try:
            return octeti.decode(enc)
        except UnicodeDecodeError:
            continue
    return ""


def citeste_docx(octeti):
    try:
        from docx import Document
    except ImportError:
        raise RuntimeError(
            "Lipsește biblioteca python-docx. Instalează-o cu "
            "'pip install python-docx' și adaug-o în requirements.txt."
        )
    doc = Document(io.BytesIO(octeti))
    parti = [p.text for p in doc.paragraphs if p.text.strip()]
    for tabel in doc.tables:
        for rand in tabel.rows:
            celule = [c.text.strip() for c in rand.cells if c.text.strip()]
            if celule:
                parti.append(" | ".join(celule))
    return "\n\n".join(parti)


def _capitole(doc):
    """Capitolele din cuprinsul PDF-ului (dacă există): titlu + interval de pagini."""
    try:
        toc = [
            (lv, t.strip(), p)
            for lv, t, p in doc.get_toc()
            if lv <= 2 and t.strip() and 1 <= p <= len(doc)
        ]
    except Exception:
        return []
    rezultat = []
    for i, (lv, titlu, p) in enumerate(toc):
        end = len(doc)
        for lv2, _, p2 in toc[i + 1 :]:
            if lv2 <= lv:
                end = max(p, p2 - 1)
                break
        rezultat.append(
            {"titlu": ("— " if lv == 2 else "") + titlu, "start": p, "end": end}
        )
    return rezultat


@st.cache_data(show_spinner="Citesc documentul...", max_entries=5)
def incarca_fisier(hash_fisier, extensie, _octeti):
    capitole = []
    if extensie == "pdf":
        doc = pymupdf.open(stream=_octeti, filetype="pdf")
        unitati = [
            {"pagina": i, "text": p.get_text()} for i, p in enumerate(doc, start=1)
        ]
        capitole = _capitole(doc)
        doc.close()
        eticheta = "pagina"
    else:
        text = citeste_docx(_octeti) if extensie == "docx" else citeste_txt(_octeti)
        unitati = imparte_text(text)
        eticheta = "secțiunea"
    return {"unitati": unitati, "capitole": capitole, "eticheta": eticheta}


def material_din_fisier(fisier):
    octeti = fisier.getvalue()
    extensie = fisier.name.rsplit(".", 1)[-1].lower()
    hash_fisier = hashlib.sha256(octeti).hexdigest()
    date = incarca_fisier(hash_fisier, extensie, octeti)
    return {**date, "hash": hash_fisier, "nume": fisier.name}


def material_din_text(text):
    return {
        "unitati": imparte_text(text),
        "capitole": [],
        "eticheta": "secțiunea",
        "hash": hashlib.sha256(("text:" + text).encode("utf-8")).hexdigest(),
        "nume": "Text lipit",
    }


def are_text(mat):
    """False dacă documentul nu conține text (de exemplu un PDF scanat)."""
    return sum(len(u["text"].strip()) for u in mat["unitati"]) >= 50


# ------------------------- alegerea părții din document -------------------------
def taie_la_buget(unitati, start, end):
    """Unitățile de la start la end (numerotate de la 1), oprite când se depășește bugetul."""
    ales, total = [], 0
    for u in unitati[start - 1 : end]:
        lungime = len(u["text"])
        if ales and total + lungime > MAX_CARACTERE:
            break
        ales.append(u)
        total += lungime
    return ales


def sari_la_capitol(cheie_interval, cheie_cap, capitole):
    i = st.session_state.get(cheie_cap)
    if i is not None:
        st.session_state[cheie_interval] = (capitole[i]["start"], capitole[i]["end"])


def alege_interval(mat):
    """Afișează alegerea părții din document. Întoarce (unități_alese, hash_material)
    sau (None, None) dacă partea aleasă nu are text."""
    unitati = mat["unitati"]
    n = len(unitati)
    plural = PLURAL[mat["eticheta"]]
    total_car = sum(len(u["text"]) for u in unitati)
    incape = total_car <= MAX_CARACTERE
    start, end = 1, n

    if n > 1:
        cheie = f"interval_{mat['hash']}"
        if cheie not in st.session_state:
            prima = taie_la_buget(unitati, 1, n)
            st.session_state[cheie] = (1, max(1, len(prima)))

        with st.expander("Ce parte din document folosesc?", expanded=not incape):
            if incape:
                st.caption(
                    "Tot documentul încape într-o singură generare. "
                    "Poți totuși să alegi doar o parte."
                )
            else:
                st.warning(
                    f"Documentul are aproximativ {fmt(total_car)} de caractere, "
                    "prea mult pentru o singură generare. Alege o parte (maximum "
                    f"{fmt(MAX_CARACTERE)} de caractere odată). Secțiunea „Întreabă” "
                    "caută oricum în tot documentul."
                )
            if mat["capitole"]:
                caps = mat["capitole"]
                cheie_cap = f"cap_{mat['hash']}"
                st.selectbox(
                    "Sari la un capitol (din cuprinsul documentului)",
                    range(len(caps)),
                    index=None,
                    format_func=lambda i: caps[i]["titlu"],
                    key=cheie_cap,
                    placeholder="Alege un capitol (opțional)",
                    on_change=sari_la_capitol,
                    args=(cheie, cheie_cap, caps),
                )
            start, end = st.slider(f"Interval ({plural})", 1, n, key=cheie)

    selectate = taie_la_buget(unitati, start, end)
    caractere = sum(len(u["text"]) for u in selectate)
    if sum(len(u["text"].strip()) for u in selectate) < 50:
        st.error(f"În {plural} {start}-{end} nu am găsit text. Alege alt interval.")
        return None, None

    ultimul = selectate[-1]["pagina"]
    if ultimul < end:
        st.warning(
            f"Intervalul ales e prea mare pentru o singură generare, așa că folosesc "
            f"doar {plural} {start}-{ultimul}. Alege un interval mai mic ca să "
            "acoperi și restul."
        )
    elif n > 1:
        st.caption(
            f"Pentru quiz, examen, flashcarduri și rezumat folosesc {plural} "
            f"{start}-{end} (aprox. {fmt(caractere)} de caractere)."
        )

    hash_material = hashlib.sha256(
        f"{mat['hash']}|{start}|{ultimul}".encode("utf-8")
    ).hexdigest()
    return selectate, hash_material


# ------------------------- căutare în tot documentul (pentru "Întreabă") -------------------------
STOP_CUVINTE = {
    "care", "este", "sunt", "cum", "cand", "unde", "cine", "pentru", "din", "dar",
    "sau", "mai", "prin", "fara", "despre", "dupa", "intre", "cat", "cati", "cate",
    "era", "tot", "toate", "iar", "nici", "doar", "asa", "foarte", "pot", "fiind",
    "acest", "aceasta", "aceste", "acesta", "unui", "unei", "fost", "poate",
    "explica", "explicati", "spune", "spuneti", "simplu", "inseamna", "ceea",
    "dintre", "are", "avea", "lui", "cei", "cea", "cel", "the", "and", "what",
    "how", "why", "who", "does", "was", "were", "with", "for", "that", "this",
    "from", "about", "which", "when", "where", "into", "explain", "tell", "mean",
    "means", "you", "your",
}


def _fara_diacritice(t):
    t = unicodedata.normalize("NFD", t.lower())
    return "".join(c for c in t if unicodedata.category(c) != "Mn")


def cuvinte_cheie(t):
    """Cuvinte fără diacritice, tăiate la 6 litere (ca "cloroplast" și
    "cloroplaste" să se potrivească)."""
    rezultat = []
    for w in re.findall(r"[a-z0-9]{3,}", _fara_diacritice(t)):
        if w not in STOP_CUVINTE:
            rezultat.append(w[:6])
    return rezultat


@st.cache_resource(max_entries=3)
def construieste_index(hash_doc, _unitati):
    return [Counter(cuvinte_cheie(u["text"])) for u in _unitati]


def alege_fragmente(mat, intrebare, intrebare_anterioara=""):
    """Întoarce (fragmente, e_partial). Dacă documentul încape în buget, îl trimite
    întreg; altfel alege paginile cele mai relevante pentru întrebare."""
    unitati = mat["unitati"]
    total = sum(len(u["text"]) for u in unitati)
    if total <= MAX_CARACTERE:
        return unitati, False

    n = len(unitati)
    index = construieste_index(mat["hash"], unitati)
    termeni = set(cuvinte_cheie(f"{intrebare} {intrebare_anterioara}"))

    df = Counter()
    for c in index:
        df.update(c.keys())

    scoruri = []
    for i, c in enumerate(index):
        s = 0.0
        for w in termeni:
            tf = c.get(w, 0)
            if tf:
                s += math.log(1 + n / df[w]) * (tf * 2.5) / (tf + 1.5)
        if s > 0:
            scoruri.append((s, i))
    scoruri.sort(reverse=True)

    ordine = [i for _, i in scoruri[:NR_FRAGMENTE_INTREBARE]]
    if not ordine:
        ordine = list(range(min(n, NR_FRAGMENTE_INTREBARE)))

    alese = []
    folosit = 0

    def adauga(i):
        nonlocal folosit
        if i < 0 or i >= n or i in alese:
            return
        lungime = len(unitati[i]["text"])
        if alese and folosit + lungime > MAX_CARACTERE:
            return
        alese.append(i)
        folosit += lungime

    for i in ordine:
        adauga(i)
    for i in ordine:  # vecinii paginilor găsite, cât ne permite bugetul
        adauga(i + 1)
        adauga(i - 1)

    alese.sort()
    return [unitati[i] for i in alese], True


def descrie_surse(fragmente, eticheta):
    nr = [str(u["pagina"]) for u in fragmente]
    if len(nr) > 15:
        nr = nr[:15] + ["..."]
    return f"Am căutat în {PLURAL[eticheta]}: " + ", ".join(nr)
